"""THETIS-MRV public data ingestion — the free lookup demo's data source.

EMSA publishes annual verified emissions per ship. Two loaders:
- load_thetis_xlsx: the REAL public export ("EU MRV Publication of
  information", XLSX, ~16k ships/year). Downloaded 8 Jul 2026 via the public
  API: /api/public-emission-report/reporting-period-document/binary/{rp}/{v}
  (list versions at /api/public-emission-report/downloadable-files).
  Headers verified against 2025-v6 and 2024-v228 files.
- load_thetis_csv: legacy simple-CSV loader kept for small fixtures/tests.

Reconstruction logic: the real export gives per-scope CO2 splits (intra-MS
voyages, departed-from-MS, arrived-to-MS, at berth) plus average CO2 per
distance, so we rebuild one synthetic leg per scope with HFO-equivalent
tonnage. Still labelled ESTIMATE — full accuracy needs owner data.
"""
import csv
import os
from dataclasses import dataclass
from typing import Optional
from .schema import (Vessel, AnnualActivity, VoyageLeg, FuelConsumption,
                     FuelType, ShipType, VoyageScope)
from .cii import CF_TCO2_PER_T

# Legacy simple-CSV layout (fixtures/tests only).
COLS = {
    "imo": "IMO Number",
    "name": "Name",
    "type": "Ship type",
    "co2_total": "Total CO₂ emissions [m tonnes]",
    "fuel_total": "Total fuel consumption [m tonnes]",
    "co2_per_dist": "Annual average CO₂ emissions per distance [kg CO₂ / n mile]",
}

# Real XLSX export: header row is row 3 (rows 1-2 are group banners); the
# ship block comes first, so first-occurrence wins for duplicated names
# ("IMO Number"/"Name" repeat in the Company block). Verified vs 2025-v6.
XLSX_HEADER_ROW = 3
XLSX_COLS = {
    "imo": "IMO Number",
    "name": "Name",
    "type": "Ship type",
    "period": "Reporting Period",
    "tech_eff": "Technical efficiency",   # e.g. "EEXI (5.34 gCO₂/t·nm)"
    "fuel_total": "Total fuel consumption [m tonnes]",
    "co2_total": "Total CO₂ emissions [m tonnes]",
    "co2_intra": "CO₂ emissions from all voyages between ports under a MS jurisdiction [m tonnes]",
    "co2_departed": "CO₂ emissions from all voyages which departed from ports under a MS jurisdiction [m tonnes]",
    "co2_arrived": "CO₂ emissions from all voyages to ports under a MS jurisdiction [m tonnes]",
    "co2_berth": "CO₂ emissions which occurred within ports under a MS jurisdiction at berth [m tonnes]",
    "co2_ets": "CO2 emissions to be reported under Directive 2003/87/EC [m tonnes]",
    "co2eq_total": "Total CO₂eq emissions [m tonnes]",
    "co2_per_dist": "CO₂ emissions per distance [kg CO₂ / n mile]",
    "fuel_per_dist": "Fuel consumption per distance [kg / n mile]",
    "time_at_sea_h": "Time spent at sea [hours]",
}

SHIPTYPE_MAP = {
    "bulk carrier": ShipType.BULK_CARRIER,
    "oil tanker": ShipType.TANKER,
    "chemical tanker": ShipType.TANKER,
    "container ship": ShipType.CONTAINER,
    "general cargo ship": ShipType.GENERAL_CARGO,
    "gas carrier": ShipType.GAS_CARRIER,
    "refrigerated cargo carrier": ShipType.REFRIGERATED_CARGO,
    "combination carrier": ShipType.COMBINATION_CARRIER,
    "lng carrier": ShipType.LNG_CARRIER,
    "ro-ro ship": ShipType.RO_RO,
    "ro-pax ship": ShipType.RO_PAX,
    "passenger ship (cruise passenger ship)": ShipType.CRUISE,
    "vehicle carrier": ShipType.VEHICLE_CARRIER,
    # Still unmapped (no G2/G4 line): "container/ro-ro cargo ship",
    # "passenger ship", "other ship types", "other ship types (offshore)"
}


@dataclass
class ThetisRecord:
    imo: str
    name: str
    ship_type: Optional[ShipType]
    ship_type_raw: str
    co2_tonnes: float
    fuel_tonnes: float
    est_distance_nm: Optional[float]
    # Real-export extras (None when loaded from the legacy CSV):
    co2_intra_t: Optional[float] = None
    co2_departed_t: Optional[float] = None
    co2_arrived_t: Optional[float] = None
    co2_berth_t: Optional[float] = None
    co2_ets_reported_t: Optional[float] = None   # verifier-computed ETS scope
    co2eq_tonnes: Optional[float] = None
    reporting_period: Optional[int] = None
    tech_efficiency: Optional[str] = None        # e.g. "EEXI (5.34 gCO₂/t·nm)"
    company: Optional[str] = None                # DoC holder (fleet grouping)


def load_thetis_xlsx(path: str, sheet: Optional[str] = None) -> dict[str, ThetisRecord]:
    """Parse the real EMSA public XLSX export into {imo: record}.

    Defaults to the first sheet ('<year> Full ERs'); pass the partial-ER
    sheet name explicitly to load partial-year reports.
    """
    import openpyxl, warnings
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")   # workbook ships without default style
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
        rows = ws.iter_rows(min_row=XLSX_HEADER_ROW, values_only=True)
        headers = [str(h).strip() if h is not None else "" for h in next(rows)]
        idx: dict[str, int] = {}
        for key, colname in XLSX_COLS.items():
            idx[key] = headers.index(colname)   # first occurrence = ship block
        # Company name = second "Name" column (company block follows ship block)
        name_cols = [i for i, h in enumerate(headers) if h == "Name"]
        company_idx = name_cols[1] if len(name_cols) > 1 else None

        out: dict[str, ThetisRecord] = {}
        for row in rows:
            imo = row[idx["imo"]]
            if imo is None:
                continue
            imo = str(imo).strip().split(".")[0]
            raw_type = str(row[idx["type"]] or "").strip()
            co2 = _f(row[idx["co2_total"]]) or 0.0
            cpd = _f(row[idx["co2_per_dist"]])
            dist = (co2 * 1000.0 / cpd) if (co2 and cpd) else None  # t->kg / (kg/nm)
            period = _f(row[idx["period"]])
            out[imo] = ThetisRecord(
                imo=imo,
                name=str(row[idx["name"]] or "").strip(),
                ship_type=SHIPTYPE_MAP.get(raw_type.lower()),
                ship_type_raw=raw_type,
                co2_tonnes=co2,
                fuel_tonnes=_f(row[idx["fuel_total"]]) or 0.0,
                est_distance_nm=dist,
                co2_intra_t=_f(row[idx["co2_intra"]]),
                co2_departed_t=_f(row[idx["co2_departed"]]),
                co2_arrived_t=_f(row[idx["co2_arrived"]]),
                co2_berth_t=_f(row[idx["co2_berth"]]),
                co2_ets_reported_t=_f(row[idx["co2_ets"]]),
                co2eq_tonnes=_f(row[idx["co2eq_total"]]),
                reporting_period=int(period) if period else None,
                tech_efficiency=(str(row[idx["tech_eff"]]).strip()
                                 if row[idx["tech_eff"]] else None),
                company=(str(row[company_idx]).strip()
                         if company_idx is not None and row[company_idx] else None),
            )
        return out
    finally:
        wb.close()


_CACHE_VERSION = 5   # bump when ThetisRecord fields or parsing change


def load_thetis_cached(path: str, sheet: Optional[str] = None) -> dict[str, ThetisRecord]:
    """load_thetis_xlsx with a pickle cache beside the file.

    The xlsx parse takes ~30 s for a 16.7k-ship export; the pickle loads in
    well under a second.

    The cache is keyed on a hash of the file's CONTENT, not its size/mtime.
    mtime changes on every `git clone`/deploy, which would force the slow
    parse on every release; a content hash only changes when EMSA's data does.
    Reading and hashing 11 MB costs a few milliseconds.

    Any failure to read the cache — missing, corrupt, written by a different
    Python version — falls back to parsing the xlsx. It must never take the
    site down.
    """
    import hashlib
    import pickle
    with open(path, "rb") as f:
        digest = hashlib.sha1(f.read()).hexdigest()
    key = (digest, sheet, _CACHE_VERSION)
    cache = path + ".pkl"
    try:
        with open(cache, "rb") as f:
            cached_key, recs = pickle.load(f)
        if cached_key == key:
            return recs
    except Exception:                                       # noqa: BLE001
        pass
    recs = load_thetis_xlsx(path, sheet)
    try:
        with open(cache, "wb") as f:
            pickle.dump((key, recs), f, protocol=4)   # readable by Python 3.4+
    except OSError:
        pass   # cache is best-effort (read-only filesystem etc.)
    return recs


def load_thetis_csv(path: str) -> dict[str, ThetisRecord]:
    """Parse the legacy simple-CSV fixture layout into {imo: record}."""
    out: dict[str, ThetisRecord] = {}
    with open(path, newline="", encoding="utf-8-sig") as f:
        for row in csv.DictReader(f):
            imo = str(row.get(COLS["imo"], "")).strip()
            if not imo:
                continue
            raw_type = str(row.get(COLS["type"], "")).strip()
            co2 = _f(row.get(COLS["co2_total"]))
            fuel = _f(row.get(COLS["fuel_total"]))
            cpd = _f(row.get(COLS["co2_per_dist"]))
            dist = (co2 * 1000.0 / cpd) if (co2 and cpd) else None  # t->kg / (kg/nm)
            out[imo] = ThetisRecord(imo, row.get(COLS["name"], "").strip(),
                                    SHIPTYPE_MAP.get(raw_type.lower()), raw_type,
                                    co2 or 0.0, fuel or 0.0, dist)
    return out


def assumed_fuel(rec: ThetisRecord) -> FuelType:
    """Which fuel the public-data reconstruction should assume.

    The export publishes CO2 and total fuel mass but NOT the fuel type, so a
    reconstruction has to assume one. Assuming HFO for everything is the
    conservative default — but it is plainly wrong for LNG carriers, which
    run on cargo boil-off gas, and it has a systematic consequence: HFO's
    WtW intensity (91.74) sits above every FuelEU target, so an all-HFO
    assumption makes *every ship in the register* look like a FuelEU deficit
    and hides the surplus side of the pooling market entirely.

    We therefore infer LNG for LNG carriers. The inference is surfaced in the
    leg description so it is never silently applied, and any owner-supplied
    bunker data overrides it.

    A reconstructed CO2-per-tonne cross-check could refine this further
    (LNG's CF 2.750 vs HFO's 3.114), but the export's "total fuel" column is
    not always populated — VERIFY before relying on it.
    """
    return FuelType.LNG if rec.ship_type == ShipType.LNG_CARRIER else FuelType.HFO


def record_to_activity(rec: ThetisRecord, year: int, dwt: float,
                       assume_fuel: Optional[FuelType] = None) -> AnnualActivity:
    """Build an ESTIMATE activity from public MRV data.

    With the real export's scope splits we rebuild one leg per jurisdiction
    scope (intra-EEA voyages, extra-EEA voyages, at berth, out of scope);
    the legacy CSV path falls back to a single intra-EEA leg. Fuel mass per
    leg is backed out of that scope's CO2 using the assumed fuel's CF
    (see `assumed_fuel`); pass `assume_fuel` to override.
    """
    if rec.ship_type is None:
        raise ValueError(f"Unmapped ship type {rec.ship_type_raw!r} for IMO "
                         f"{rec.imo} — extend SHIPTYPE_MAP (and G2/G4 tables).")
    fuel = assume_fuel or assumed_fuel(rec)
    cf = CF_TCO2_PER_T[fuel]
    tag = f"assumed {fuel.value.upper()}"
    # `dwt` is the caller-supplied capacity: DWT for most types, GT for the
    # GT-capacity types (ro-ro/ro-pax/cruise/vehicle) — set both so either
    # basis works; the CII engine picks the right one per MEPC.353(78).
    vessel = Vessel(rec.imo, rec.name, rec.ship_type, dwt=dwt, gt=dwt)
    total_dist = rec.est_distance_nm or 0.0
    total_co2 = rec.co2_tonnes or 0.0

    if rec.co2_intra_t is None:   # legacy CSV: no splits available
        fuel_t = rec.fuel_tonnes or (total_co2 / cf)
        leg = VoyageLeg(VoyageScope.INTRA_EEA, total_dist,
                        [FuelConsumption(fuel, fuel_t)],
                        f"THETIS-MRV verified EU scope (ESTIMATE, {tag})")
        return AnnualActivity(vessel, year, [leg], fuel_is_combusted=True)

    intra = rec.co2_intra_t or 0.0
    extra = (rec.co2_departed_t or 0.0) + (rec.co2_arrived_t or 0.0)
    berth = rec.co2_berth_t or 0.0
    out_of_scope = max(0.0, total_co2 - intra - extra - berth)

    legs: list[VoyageLeg] = []
    # The register's distance (total CO2 / CO2 per mile) is the whole year's distance. Berth
    # time burns fuel without sailing, so all of it goes to the sea legs, by their CO2 share.
    sea_co2 = intra + extra + out_of_scope
    for scope, co2_t, label in (
            (VoyageScope.INTRA_EEA, intra, "intra-EEA voyages"),
            (VoyageScope.EXTRA_EEA, extra, "voyages to/from EEA"),
            (VoyageScope.AT_BERTH_EEA, berth, "at berth in EEA"),
            (VoyageScope.OUT_OF_SCOPE, out_of_scope, "outside EEA scope")):
        if co2_t <= 0:
            continue
        dist = 0.0 if scope == VoyageScope.AT_BERTH_EEA else (total_dist * co2_t / sea_co2 if sea_co2 else 0.0)
        legs.append(VoyageLeg(scope, dist,
                              [FuelConsumption(fuel, co2_t / cf)],
                              f"THETIS-MRV {label} (ESTIMATE, {tag})"))
    return AnnualActivity(vessel, year, legs, ets_calibration=ets_calibration(rec),
                          fuel_is_combusted=True)


def rebuilt_ets_co2(rec: "ThetisRecord") -> Optional[float]:
    """Our rebuild of the ETS-scope CO2 from the register's scope split: voyages inside the
    EU 100%, to or from the EU 50%, at berth 100% (Directive Art. 3ga). None without a split."""
    if rec.co2_intra_t is None:
        return None
    return ((rec.co2_intra_t or 0.0)
            + 0.5 * ((rec.co2_departed_t or 0.0) + (rec.co2_arrived_t or 0.0))
            + (rec.co2_berth_t or 0.0))


def ets_calibration(rec: "ThetisRecord") -> Optional[float]:
    """Verifier-reported ETS-scope CO2 divided by our rebuild of it, or None when the register
    gives no reported figure or our rebuild is zero (so there is nothing to scale)."""
    rebuilt = rebuilt_ets_co2(rec)
    reported = rec.co2_ets_reported_t
    if reported is None or not rebuilt:
        return None
    return reported / rebuilt


def _f(x) -> Optional[float]:
    if x is None:
        return None
    try:
        return float(str(x).replace(",", "").strip())
    except (TypeError, ValueError):
        return None   # handles 'N/A', 'Division by zero!', blanks
