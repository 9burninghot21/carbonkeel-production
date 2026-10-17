#!/usr/bin/env python3
"""Check for, and fetch, EMSA's public THETIS-MRV publication files.

Source: https://mrv.emsa.europa.eu/  (the "GHG Emission Report" page).
EMSA publishes one spreadsheet per reporting year, which it regenerates as
operators and verifiers submit amendments (the version number rises). The
files are public by law (Article 21 of Regulation (EU) 2015/757); see
docs/DATA_SOURCES.md for the licence position and the attribution we owe.

  python scripts/fetch_thetis.py                      check only (metadata, no download)
  python scripts/fetch_thetis.py --download 2025 --yes   fetch the newest 2025 file
  python scripts/fetch_thetis.py --download 2025 2024 --yes

Design rules
  * The default does not download anything. A download needs --yes.
  * One request at a time, a pause between requests, an honest User-Agent with
    a contact address, and no retry storms. We take only what we use.
  * Old files are never deleted; a new version is saved next to the old one.
  * A download is checked before it is kept: it must be a real .xlsx, have the
    expected sheets and columns, and its totals are compared with the file it
    would replace so a silent format change cannot reach the app.

Exit codes: 0 up to date / download kept, 3 newer files exist (check mode),
1 error or a failed validation.
"""
import argparse
import hashlib
import io
import json
import os
import sys
import time
import urllib.error
import urllib.request
import warnings
import zipfile
from datetime import datetime, timezone

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
warnings.filterwarnings("ignore")

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
sys.path.insert(0, os.path.join(ROOT, "src"))
from ghgcore.exports import newest_per_year, parse  # noqa: E402

BASE = "https://mrv.emsa.europa.eu/api/public-emission-report"
LIST_URL = f"{BASE}/downloadable-files"
BIN_URL = f"{BASE}/reporting-period-document/binary/{{year}}/{{version}}"
FIXTURES = os.path.join(ROOT, "fixtures")
MANIFEST = os.path.join(ROOT, "data", "source_manifest.json")
UA = "CarbonKeel-data-fetch/1.0 (+mailto:enquiry@carbonkeel.com; public MRV data, low volume)"
PAUSE_S = 2.0


def _open(url, method="GET", timeout=60):
    req = urllib.request.Request(url, method=method, headers={"User-Agent": UA})
    last = None
    for attempt in range(3):
        try:
            return urllib.request.urlopen(req, timeout=timeout)
        except urllib.error.HTTPError as e:
            if e.code in (429, 503) and attempt < 2:      # asked to slow down: do so
                time.sleep(10 * (attempt + 1))
                last = e
                continue
            raise
        except (urllib.error.URLError, TimeoutError) as e:
            last = e
            time.sleep(3 * (attempt + 1))
    raise last


def live_listing():
    with _open(LIST_URL) as r:
        data = json.load(r)
    rows = []
    for x in data.get("results", []):
        rows.append({"year": int(x["reportingPeriod"]), "version": int(x["version"]),
                     "generated": x.get("generationDate", ""), "file": x.get("fileName", "")})
    return sorted(rows, key=lambda r: r["year"])


def remote_size(year, version):
    time.sleep(PAUSE_S)
    with _open(BIN_URL.format(year=year, version=version), method="HEAD") as r:
        return int(r.headers.get("Content-Length") or 0)


def check():
    local = newest_per_year(FIXTURES)
    rows = live_listing()
    print(f"EMSA THETIS-MRV public listing ({len(rows)} reporting periods), checked "
          f"{datetime.now(timezone.utc):%d %b %Y %H:%M} UTC\n")
    print(f"{'year':>5}  {'we hold':>10}  {'EMSA has':>9}  {'generated':<20} status")
    newer = 0
    for r in rows:
        have = parse(local[r["year"]])[1] if r["year"] in local else None
        if have is None:
            status = "not held"
        elif have == r["version"]:
            status = "up to date"
        elif have < r["version"]:
            status = f"NEWER on EMSA (+{r['version'] - have} versions)"
            newer += 1
        else:
            status = "we hold a newer version than EMSA lists"
        print(f"{r['year']:>5}  {('v' + str(have)) if have else '-':>10}  v{r['version']:<8} "
              f"{r['generated']:<20} {status}")
    held_years = sorted(local)
    stale = [y for y in held_years if y in {r['year'] for r in rows}
             and parse(local[y])[1] < next(r['version'] for r in rows if r['year'] == y)]
    if stale:
        print(f"\nStale: {', '.join(map(str, stale))}. To refresh (downloads about 10 MB per file):\n"
              f"  python scripts/fetch_thetis.py --download {' '.join(map(str, stale))} --yes")
    return 3 if newer else 0


def _sheet_totals(path):
    """(ships on the first sheet, ETS-scope CO2 tonnes, sheet names, missing columns)."""
    import openpyxl
    from ghgcore.thetis import XLSX_COLS, XLSX_HEADER_ROW
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    names = list(wb.sheetnames)
    ws = wb[names[0]]
    header = None
    ships, ets = 0, 0.0
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        if i == XLSX_HEADER_ROW:
            header = [str(c).strip() if c is not None else "" for c in row]
            idx = header.index(XLSX_COLS["co2_ets"]) if XLSX_COLS["co2_ets"] in header else None
        elif i > XLSX_HEADER_ROW and row and row[0] is not None:
            ships += 1
            if idx is not None:
                try:
                    ets += float(str(row[idx]).replace(",", ""))
                except (TypeError, ValueError):
                    pass
    wb.close()
    missing = [v for v in XLSX_COLS.values() if not header or v not in header]
    return ships, ets, names, missing


def validate(path, previous):
    """Raise ValueError with a clear reason if the file should not be used."""
    if not zipfile.is_zipfile(path):
        raise ValueError("not a valid .xlsx (zip) file")
    ships, ets, names, missing = _sheet_totals(path)
    print(f"    sheets: {names}")
    if not names or "Full ERs" not in names[0]:
        raise ValueError(f"first sheet is {names[:1]}, expected '<year> Full ERs'")
    if missing:
        raise ValueError(f"columns we rely on are missing or renamed: {missing}")
    print(f"    ships on the first sheet: {ships:,}; ETS-scope CO2: {ets / 1e6:,.2f} Mt")
    if previous:
        p_ships, p_ets, _, _ = _sheet_totals(previous)
        d_ships = (ships - p_ships) / p_ships * 100 if p_ships else 0
        d_ets = (ets - p_ets) / p_ets * 100 if p_ets else 0
        print(f"    against {os.path.basename(previous)}: ships {d_ships:+.1f}%, ETS CO2 {d_ets:+.1f}%")
        if abs(d_ships) > 25 or abs(d_ets) > 25:
            raise ValueError("totals moved by more than 25% against the file we hold; "
                             "check by hand before using it")
    return ships, ets


def _manifest_add(entry):
    os.makedirs(os.path.dirname(MANIFEST), exist_ok=True)
    data = []
    if os.path.exists(MANIFEST):
        try:
            data = json.load(open(MANIFEST, encoding="utf-8"))
        except ValueError:
            data = []
    data.append(entry)
    with open(MANIFEST, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2)


def download(years):
    rows = {r["year"]: r for r in live_listing()}
    local = newest_per_year(FIXTURES)
    for y in years:
        if y not in rows:
            print(f"{y}: EMSA does not list this reporting period"); return 1
        r = rows[y]
        dest = os.path.join(FIXTURES, f"thetis_mrv_{y}_v{r['version']}.xlsx")
        if os.path.exists(dest):
            print(f"{y}: v{r['version']} is already held ({os.path.basename(dest)})"); continue
        size = remote_size(y, r["version"])
        print(f"{y}: fetching version {r['version']} (generated {r['generated']}), "
              f"{size / 1e6:.1f} MB ...")
        time.sleep(PAUSE_S)
        part = dest + ".part"
        h = hashlib.sha256()
        got = 0
        with _open(BIN_URL.format(year=y, version=r["version"]), timeout=180) as resp, \
                open(part, "wb") as out:
            while True:
                chunk = resp.read(1 << 16)
                if not chunk:
                    break
                out.write(chunk); h.update(chunk); got += len(chunk)
        if size and got != size:
            os.remove(part); print(f"    size mismatch ({got} of {size} bytes); discarded"); return 1
        try:
            ships, ets = validate(part, local.get(y))
        except Exception as e:                           # noqa: BLE001
            os.remove(part); print(f"    REJECTED: {e}"); return 1
        os.replace(part, dest)
        _manifest_add({"reporting_year": y, "version": r["version"], "file": os.path.basename(dest),
                       "emsa_generated": r["generated"], "source": LIST_URL,
                       "retrieved_utc": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
                       "bytes": got, "sha256": h.hexdigest(), "ships_first_sheet": ships,
                       "ets_scope_co2_t": round(ets)})
        print(f"    kept as {os.path.basename(dest)}  (sha256 {h.hexdigest()[:16]}...)")
        time.sleep(PAUSE_S)
    print("\nNext: python scripts/market_data.py, then python scripts/preflight.py")
    print("(the app and scripts pick the highest version automatically; old files are kept)")
    return 0


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--download", nargs="+", type=int, metavar="YEAR",
                    help="fetch the newest file for these reporting years")
    ap.add_argument("--yes", action="store_true", help="confirm that downloading is wanted")
    a = ap.parse_args()
    try:
        if a.download:
            if not a.yes:
                print("Add --yes to confirm. This downloads about 10 MB per file from "
                      "mrv.emsa.europa.eu.")
                return 1
            return download(a.download)
        return check()
    except (urllib.error.URLError, TimeoutError, OSError) as e:
        print(f"Could not reach EMSA: {e}")
        return 1


if __name__ == "__main__":
    sys.exit(main())
