#!/usr/bin/env python3
"""Market numbers from the EMSA THETIS-MRV public register.

Everything the landing page and docs/NEXT_REQUIREMENTS.md say about market
size comes from here, so no figure is typed by hand. Output:
  data/market_stats.json      machine-readable (the landing page numbers)
  prints a readable summary

All euro figures are ESTIMATES: verifier-reported ETS-scope CO2 x an assumed
allowance price. The price is an input, not a forecast; sensitivity is shown.

Usage: python scripts/market_data.py
"""
import glob, io, json, os, sys, warnings
from collections import defaultdict

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
warnings.filterwarnings("ignore")

ROOT = os.path.join(os.path.dirname(__file__), "..")
sys.path.insert(0, os.path.join(ROOT, "src"))
from ghgcore.exports import newest_per_year
from ghgcore.assumptions import DEFAULT_EUA_PRICE_EUR
FOCUS = {"container ship", "bulk carrier", "oil tanker", "chemical tanker"}
BASE_PRICE = int(DEFAULT_EUA_PRICE_EUR)
EUA_PRICES = (70, BASE_PRICE, 100)
IMO, NAME, TYPE, CO_NAME, CO2_TOTAL, CO2_ETS = 0, 1, 2, 9, 28, 37


def _f(x):
    try:
        return float(str(x).replace(",", "").strip())
    except (TypeError, ValueError):
        return 0.0


def read(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    ships = []
    for row in ws.iter_rows(min_row=4, values_only=True):
        if row[IMO] is None:
            continue
        ships.append({
            "imo": str(row[IMO]).split(".")[0],
            "type": str(row[TYPE] or "").strip().lower(),
            "co": str(row[CO_NAME] or "").strip(),
            "co2": _f(row[CO2_TOTAL]),
            "ets": _f(row[CO2_ETS]),
        })
    wb.close()
    return ships


def main():
    per_year = newest_per_year(os.path.join(ROOT, "fixtures"))
    years = sorted(per_year)
    files = [per_year[y] for y in years]
    cur = read(files[-1])
    prev = read(files[-2]) if len(files) > 1 else []

    ets_total = sum(s["ets"] for s in cur)
    by_type = defaultdict(lambda: [0, 0.0])
    for s in cur:
        by_type[s["type"]][0] += 1
        by_type[s["type"]][1] += s["ets"]

    comp = defaultdict(lambda: {"n": 0, "focus": 0, "ets": 0.0})
    for s in cur:
        if not s["co"]:
            continue
        c = comp[s["co"]]
        c["n"] += 1
        c["focus"] += s["type"] in FOCUS
        c["ets"] += s["ets"]

    def band(lo, hi):
        cs = [c for c in comp.values() if lo <= c["n"] <= hi]
        return {"companies": len(cs), "ships": sum(c["n"] for c in cs),
                "ets_t": round(sum(c["ets"] for c in cs))}

    icp = [c for c in comp.values()
           if 10 <= c["n"] <= 50 and c["focus"] / c["n"] >= 0.5]
    ranked = sorted(comp.values(), key=lambda c: -c["ets"])
    top10pct = ranked[:max(1, len(ranked) // 10)]

    out = {
        "source": os.path.basename(files[-1]),
        "ships": len(cur),
        "companies": len(comp),
        "ets_scope_co2_mt": round(ets_total / 1e6, 1),
        "assumed_eua_eur": BASE_PRICE,
        "implied_ets_cost_eur_bn": {str(p): round(ets_total * p / 1e9, 2) for p in EUA_PRICES},
        "ships_prev_year": len(prev),
        "ets_prev_year_mt": round(sum(s["ets"] for s in prev) / 1e6, 1) if prev else None,
        "avg_ets_t_per_ship": round(ets_total / len(cur)),
        "avg_cost_per_ship_eur": round(ets_total / len(cur) * BASE_PRICE),
        "types": {t: {"ships": n, "ets_mt": round(e / 1e6, 2)}
                  for t, (n, e) in sorted(by_type.items(), key=lambda kv: -kv[1][1])[:8]},
        "size_bands": {"1": band(1, 1), "2-9": band(2, 9),
                       "10-50": band(10, 50), "51+": band(51, 10**6)},
        "icp": {"companies": len(icp), "ships": sum(c["n"] for c in icp),
                "ets_mt": round(sum(c["ets"] for c in icp) / 1e6, 1),
                "est_cost_eur_bn": round(sum(c["ets"] for c in icp) * BASE_PRICE / 1e9, 2)},
        "top10pct_companies_share_of_ets": round(
            sum(c["ets"] for c in top10pct) / ets_total, 3),
    }
    os.makedirs(os.path.join(ROOT, "data"), exist_ok=True)
    with open(os.path.join(ROOT, "data", "market_stats.json"), "w", encoding="utf-8") as f:
        json.dump(out, f, indent=2)
    print(json.dumps(out, indent=2))


if __name__ == "__main__":
    main()
